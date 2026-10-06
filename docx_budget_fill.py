"""Бюджет в Word-заявке донора.

Два способа записи (в обоих меняется только бюджетная таблица, весь остальной
текст и таблицы заявки остаются как есть):

1. Таблица статей расходов («№ | Статья | Кол-во | Стоимость» + строка «ИТОГО»),
   в том числе ВЛОЖЕННАЯ в ячейку-вопрос. Столбцы распознаются по шапке, строки
   при необходимости ДОБАВЛЯЮТСЯ (копия пустой строки шаблона) перед «ИТОГО»,
   итог считает код, а не модель.
2. Произвольная сетка (LLM сопоставляет цифры с адресами ячеек) — запасной путь.

Если вписать некуда — build_budget_docx собирает отдельный Word-документ с
бюджетом, который пользователь переносит в заявку сам.
"""

import copy
import io
import json
import logging
import re
import zipfile

import docx
from docx.oxml.ns import qn

from xlsx_patch import _PLACEHOLDER

logger = logging.getLogger("fund4pro.docx_budget_fill")

_BUDGET_WORDS = re.compile(
    r"бюджет|budget|смет|затрат|расход|cost|expens|стоимост|итого|total|сумм|amount|price|цена", re.I)
_TOTAL_LABEL = re.compile(r"^\s*(итого|всего|итог\b|total|subtotal|grand total|сумма по|общая сумма)", re.I)
_ADDR = re.compile(r"^r(\d+)c(\d+)$")
_MONEY = re.compile(r"^[\$€£]?\s*-?[\d\s ,.]+\s*[\$€£]?$")
MAX_ROWS_IN_PROMPT = 90

_ROLE_PATTERNS = [
    ("num", re.compile(r"^\s*(№|n°|no\.?|#|п/п|№\s*п/п)\s*$", re.I)),
    ("comment", re.compile(r"комментар|примечан|note|comment|remark|источник|source", re.I)),
    ("unit_cost", re.compile(r"цена|ставка|стоимость за|за единиц|unit cost|unit price|\brate\b|\bprice\b", re.I)),
    ("qty_unit", re.compile(r"единиц|количеств|кол-во|измерен|период|\bunit|quantity|\bqty\b|месяц", re.I)),
    ("name", re.compile(
        r"мероприят|статья|статьи|наименован|название|описание|вид расход|категори|activity|\bitem\b|"
        r"description|expense|budget line|line item", re.I)),
    ("amount", re.compile(r"стоимост|сумма|итого|всего|\bcost|amount|total|\bsum\b", re.I)),
]


# ---------------------------------------------------------------- таблицы

def _grid(table):
    """grid[r][c] = (cell, is_anchor). Объединённая ячейка: якорь — первая
    встреча в порядке строк/столбцов, остальные — не-якоря (писать нельзя)."""
    seen = []
    out = []
    for row in table.rows:
        line = []
        for cell in row.cells:
            tc = cell._tc
            anchor = not any(tc is s for s in seen)
            if anchor:
                seen.append(tc)
            line.append((cell, anchor))
        out.append(line)
    return out


def _all_tables(document) -> list:
    """Все таблицы документа, включая вложенные в ячейки (в порядке обхода)."""
    result = []

    def walk(tables):
        for t in tables:
            result.append(t)
            for row in _grid(t):
                for cell, anchor in row:
                    if anchor and cell.tables:
                        walk(cell.tables)

    walk(document.tables)
    return result


def _text(cell) -> str:
    return cell.text.strip().replace("\n", " ")


def table_roles(table):
    """(индекс_строки_шапки, {столбец: роль}) для таблицы статей расходов
    (есть и столбец названия статьи, и столбец суммы) или None."""
    grid = _grid(table)
    for hr in range(min(2, len(grid))):
        roles: dict[int, str] = {}
        for c, (cell, anchor) in enumerate(grid[hr]):
            if not anchor:
                continue
            t = _text(cell)
            if not t:
                continue
            for role, rx in _ROLE_PATTERNS:
                if rx.search(t):
                    roles[c] = role
                    break
        if "name" in roles.values() and "amount" in roles.values():
            amounts = [c for c, r in roles.items() if r == "amount"]
            for c in amounts[:-1]:
                roles[c] = "skip"
            return hr, roles
    return None


def _looks_budget(table, lenient: bool) -> bool:
    if table_roles(table) is not None:
        return True
    rows = len(table.rows)
    cols = len(table.columns)
    if rows < (3 if lenient else 5) or cols < 2:
        return False
    if lenient:
        return True
    texts = [" ".join(_text(c) for c in r.cells) for r in table.rows]
    head = " ".join(texts[:2])
    body = " ".join(texts)
    return bool(_BUDGET_WORDS.search(head) or (_BUDGET_WORDS.search(body) and len(re.findall(r"\d{3,}", body)) >= 2))


def find_line_item_tables(content: bytes) -> list[int]:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception:
        return []
    return [i for i, t in enumerate(_all_tables(document)) if table_roles(t) is not None]


def find_budget_tables(content: bytes, lenient: bool = False) -> list[int]:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception:
        return []
    return [i for i, t in enumerate(_all_tables(document)) if _looks_budget(t, lenient)]


def render_tables(content: bytes, indices: list[int]) -> str:
    """Сетка таблиц для LLM: адрес ячейки rNcM (строка, столбец с 0), ∅ — пусто
    (можно писать), ▒ — внутри объединённой, текст — подписи донора."""
    document = docx.Document(io.BytesIO(content))
    tables = _all_tables(document)
    lines = []
    for ti in indices:
        table = tables[ti]
        grid = _grid(table)
        lines.append(f"[Таблица T{ti}: строк {len(grid)}, столбцов {len(table.columns)}]")
        for r, row in enumerate(grid[:MAX_ROWS_IN_PROMPT]):
            parts = []
            for c, (cell, anchor) in enumerate(row):
                if not anchor:
                    parts.append(f"r{r}c{c}=▒")
                else:
                    t = _text(cell)
                    parts.append(f"r{r}c{c}={t[:50]!r}" if t else f"r{r}c{c}=∅")
            lines.append(" | ".join(parts))
    return "\n".join(lines)


# ------------------------------------------------- структура бюджета (LLM)

def _num(text: str):
    t = (text or "").strip() if isinstance(text, str) else str(text)
    if not t or not _MONEY.match(t):
        return None
    t = re.sub(r"[\$€£\s ]", "", t)
    t = re.sub(r"(?<=\d),(?=\d{3}(?!\d))", "", t).replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def fmt_amount(x: float) -> str:
    if abs(x - round(x)) < 0.005:
        return f"{int(round(x)):,}".replace(",", " ")
    return f"{x:,.2f}".replace(",", " ")


async def extract_budget_struct(budget_text: str) -> dict | None:
    """Согласованный бюджет (проза/markdown) -> {"currency", "items": [{category,
    name, qty, unit_cost, amount}], "total"}. Итог считает код по строкам."""
    from llm import call_claude

    system_prompt = (
        "Из согласованного бюджета проекта выдели ТОЛЬКО строки расходов (без "
        "комментариев, вопросов пользователю, итогов по категориям и общего итога) "
        "и верни СТРОГО JSON без пояснений:\n"
        '{"currency": "USD", "items": [{"category": "Зарплаты", "name": "Руководитель проекта '
        '(50% ставки)", "qty": "12 мес. × 50%", "unit_cost": 1000, "amount": 6000}]}\n'
        "Правила: одна статья бюджета — одна строка; amount — итог строки числом (как в "
        "бюджете, без валютных символов и разделителей тысяч); unit_cost — цена за единицу "
        "числом или null; qty — короткая запись количества/расчёта («3 дня × 25 чел.»); "
        "category — укрупнённая категория (Зарплаты, Админ-расходы, Мероприятия, Гонорары, "
        "Публикации, Оборудование, Банковские расходы и т.п.). Названия — на языке бюджета. "
        "Ничего не придумывай и не пересчитывай: только то, что есть в бюджете."
    )
    try:
        raw = await call_claude(system_prompt, budget_text, max_tokens=4000)
        data = json.loads(raw[raw.find("{"):raw.rfind("}") + 1])
    except Exception as exc:
        logger.warning("extract_budget_struct failed: %s: %s", type(exc).__name__, exc)
        return None
    items = []
    for it in data.get("items") or []:
        if not isinstance(it, dict):
            continue
        amount = _num(str(it.get("amount"))) if it.get("amount") is not None else None
        name = str(it.get("name") or "").strip()
        if amount is None or not name:
            continue
        unit_cost = _num(str(it.get("unit_cost"))) if it.get("unit_cost") is not None else None
        items.append({
            "category": str(it.get("category") or "").strip(),
            "name": name,
            "qty": str(it.get("qty") or "").strip(),
            "unit_cost": unit_cost,
            "amount": amount,
        })
    if not items:
        return None
    return {"currency": str(data.get("currency") or "").strip(), "items": items,
            "total": sum(i["amount"] for i in items)}


def _entries(struct: dict) -> list[dict]:
    """Строки для записи: заголовки категорий (если их >1) и статьи с номерами."""
    categories = [i["category"] for i in struct["items"]]
    use_categories = len({c for c in categories if c}) > 1
    out, last, n = [], None, 0
    for it in struct["items"]:
        if use_categories and it["category"] and it["category"] != last:
            out.append({"kind": "category", "name": it["category"]})
            last = it["category"]
        n += 1
        out.append({"kind": "item", "num": str(n), **it})
    return out


# ------------------------------------------------------- запись в таблицу

def _write(cell, text: str, bold: bool = False) -> None:
    text = str(text).replace("\n", " ").strip()
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    if first.runs:
        run = first.runs[0]
        run.text = text
        for extra in first.runs[1:]:
            extra.text = ""
    else:
        run = first.add_run(text)
        ppr = first._p.find(qn("w:pPr"))
        mark = ppr.find(qn("w:rPr")) if ppr is not None else None
        if mark is not None:
            run._r.insert(0, copy.deepcopy(mark))
    if bold:
        run.bold = True
    for p in paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def _blank_clone(tr):
    new = copy.deepcopy(tr)
    for t in new.iter(qn("w:t")):
        t.text = ""
    for br in list(new.iter(qn("w:br"))):
        br.getparent().remove(br)
    return new


def _writable(cell) -> bool:
    t = _text(cell)
    return not t or bool(_PLACEHOLDER.match(t)) or bool(re.fullmatch(r"\d+[.)]?", t))


def fill_line_item_table(table, header_row: int, roles: dict, struct: dict) -> dict | None:
    """Записывает статьи построчно, добавляя строки перед «ИТОГО»; None, если
    в таблице нет строки-образца."""
    grid = _grid(table)
    total_r = next((r for r in range(len(grid) - 1, header_row, -1)
                    if any(a and _TOTAL_LABEL.match(_text(c)) for c, a in grid[r])), None)
    end = total_r if total_r is not None else len(grid)
    data_rows = list(range(header_row + 1, end))
    if not data_rows:
        return None
    entries = _entries(struct)
    template_tr = table.rows[data_rows[0]]._tr
    ref_tr = table.rows[total_r]._tr if total_r is not None else table.rows[-1]._tr
    added = max(0, len(entries) - len(data_rows))
    for _ in range(added):
        new = _blank_clone(template_tr)
        if total_r is not None:
            ref_tr.addprevious(new)
        else:
            ref_tr.addnext(new)
            ref_tr = new

    grid = _grid(table)
    col_of = {}
    for c, role in roles.items():
        col_of.setdefault(role, c)
    name_col = col_of["name"]
    written = 0
    for i, entry in enumerate(entries):
        r = header_row + 1 + i
        row = grid[r]

        def put(role, value, bold=False):
            c = col_of.get(role)
            if c is None or value in (None, ""):
                return
            cell, anchor = row[c]
            if anchor and _writable(cell):
                _write(cell, value, bold)

        if entry["kind"] == "category":
            put("name", entry["name"], bold=True)
        else:
            put("num", entry["num"])
            put("name", entry["name"])
            put("qty_unit", entry["qty"])
            if entry.get("unit_cost") is not None:
                put("unit_cost", fmt_amount(entry["unit_cost"]))
            put("amount", fmt_amount(entry["amount"]))
            written += 1
    if total_r is not None:
        total_row = _grid(table)[total_r + added]
        c = col_of["amount"]
        cell, anchor = total_row[c]
        if anchor and _writable(cell):
            _write(cell, fmt_amount(struct["total"]), bold=True)
    return {"items": written, "rows_added": added, "total": struct["total"]}


# ------------------------------------------------- запасной путь: сетка (LLM)

async def map_budget_to_tables(grid_text: str, budget_text: str) -> dict:
    from llm import call_claude

    system_prompt = (
        "Перед тобой таблицы Word-заявки донора (по строкам; у каждой ячейки адрес "
        "rNcM: N — строка, M — столбец, оба с нуля) и согласованный с пользователем "
        "бюджет. Знаки: ∅ — пустая ячейка (сюда можно писать), ▒ — внутри "
        "объединённой (не трогать), текст в кавычках — заголовки/подписи донора "
        "(не менять; подсказки вроде «укажите…» можно заменить).\n\n"
        "Найди таблицу бюджета и распиши бюджет по её строкам: определи по шапке "
        "смысл столбцов (статья, количество, ставка, период, сумма, комментарий), "
        "впиши по одной статье бюджета на строку подряд. Итоговые суммы строк и "
        "блоков («Итого») тоже впиши — рассчитай аккуратно. Не придумывай статьи "
        "и суммы: только то, что есть в согласованном бюджете; строки таблицы, "
        "которым нет соответствия, оставь пустыми. Если пустых строк меньше, чем "
        "статей, объедини мелкие статьи, не выходя за таблицу. Числа — без "
        "валютных символов и разделителей тысяч, десятичная точка. Не уверен — "
        "пропусти ячейку.\n\n"
        'Верни СТРОГО JSON без пояснений: {"T0": {"r3c1": "Руководитель проекта", "r3c4": 12000}}'
    )
    try:
        raw = await call_claude(system_prompt, f"Таблицы заявки:\n{grid_text}\n\nСогласованный бюджет:\n{budget_text}", max_tokens=4000)
        start, end = raw.find("{"), raw.rfind("}")
        data = json.loads(raw[start:end + 1])
    except Exception as exc:
        logger.warning("map_budget_to_tables failed: %s: %s", type(exc).__name__, exc)
        return {}
    return data if isinstance(data, dict) else {}


def _fmt(x: float) -> str:
    return str(int(round(x))) if abs(x - round(x)) < 0.005 else f"{x:.2f}"


def _save_faithful(document, original: bytes) -> bytes:
    """Сохраняет документ, но в исходный архив подменяет ТОЛЬКО основную часть
    (document.xml): стили, настройки, нумерация и прочее остаются байт-в-байт
    (python-docx при обычном save() переписывает их все)."""
    part_name = str(document.part.partname).lstrip("/")
    buf = io.BytesIO()
    document.save(buf)
    try:
        new_main = zipfile.ZipFile(io.BytesIO(buf.getvalue())).read(part_name)
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, new_main if item.filename == part_name else zin.read(item.filename))
        return out.getvalue()
    except Exception:
        logger.warning("_save_faithful failed, returning python-docx output", exc_info=True)
        return buf.getvalue()


def fill_docx_budget(content: bytes, mapping: dict):
    """Запись по адресам ячеек. Возвращает (файл, записанные ['T0!r3c1'],
    пропущенные [(адрес, причина)], заменённые_подсказки, пересчитанные_итоги)."""
    document = docx.Document(io.BytesIO(content))
    tables = _all_tables(document)
    applied, skipped, replaced, fixed = [], [], [], []
    written: dict[tuple[int, int, int], float | None] = {}

    for tkey, cells in mapping.items():
        digits = re.sub(r"\D", "", str(tkey))
        if not digits or not isinstance(cells, dict) or int(digits) >= len(tables):
            skipped.append((str(tkey), "таблица не найдена"))
            continue
        ti = int(digits)
        grid = _grid(tables[ti])
        for addr, value in cells.items():
            label = f"T{ti}!{addr}"
            m = _ADDR.match(str(addr).strip().lower())
            if not m or value is None or str(value).strip() == "":
                continue
            r, c = int(m.group(1)), int(m.group(2))
            if r >= len(grid) or c >= len(grid[r]):
                skipped.append((label, "вне таблицы"))
                continue
            cell, anchor = grid[r][c]
            if not anchor:
                skipped.append((label, "часть объединённой ячейки"))
                continue
            existing = _text(cell)
            was_placeholder = False
            if existing:
                if not _PLACEHOLDER.match(existing):
                    skipped.append((label, "ячейка уже заполнена текстом донора"))
                    continue
                was_placeholder = True
            _write(cell, value)
            applied.append(label)
            if was_placeholder:
                replaced.append(label)
            written[(ti, r, c)] = _num(str(value))

    for ti in sorted({k[0] for k in written}):
        grid = _grid(tables[ti])
        total_rows = [r for r, row in enumerate(grid)
                      if any(_TOTAL_LABEL.match(_text(cell)) for cell, anchor in row if anchor)]
        for r in total_rows:
            for c in range(len(grid[r])):
                if (ti, r, c) not in written or written[(ti, r, c)] is None:
                    continue
                total = 0.0
                for rr in range(1, r):
                    if rr in total_rows:
                        continue
                    cell, anchor = grid[rr][c]
                    if anchor and (v := _num(_text(cell))) is not None:
                        total += v
                llm_value = written[(ti, r, c)]
                if total > 0 and abs(total - llm_value) > max(0.5, 0.005 * total):
                    _write(grid[r][c][0], _fmt(total))
                    fixed.append(f"T{ti}!r{r}c{c}: {_fmt(llm_value)} -> {_fmt(total)}")

    return _save_faithful(document, content), applied, skipped, replaced, fixed


def fill_budget_into_docx(content: bytes, struct: dict):
    """Пытается вписать структурированный бюджет в таблицу статей расходов
    (в том числе вложенную). Возвращает (файл, отчёт) или None."""
    document = docx.Document(io.BytesIO(content))
    for ti, table in enumerate(_all_tables(document)):
        found = table_roles(table)
        if not found:
            continue
        header_row, roles = found
        report = fill_line_item_table(table, header_row, roles, struct)
        if report:
            report["table"] = ti
            return _save_faithful(document, content), report
    return None


# ------------------------------------------------- отдельный Word-документ

def build_budget_docx(struct: dict | None, budget_text: str, title: str, notes: list[str]) -> bytes:
    """Отдельный документ с бюджетом (таблица «№ | Статья | Расчёт | Сумма»);
    без структуры — текст бюджета как есть."""
    document = docx.Document()
    document.add_heading(title, level=1)
    for n in notes:
        document.add_paragraph(n)
    if struct:
        cur = struct.get("currency") or ""
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for c, h in enumerate(["№", "Статья расходов", "Расчёт (кол-во, ставка)", f"Сумма{', ' + cur if cur else ''}"]):
            table.rows[0].cells[c].text = h
            for r in table.rows[0].cells[c].paragraphs[0].runs:
                r.bold = True
        for e in _entries(struct):
            row = table.add_row().cells
            if e["kind"] == "category":
                row[1].text = e["name"]
                for r in row[1].paragraphs[0].runs:
                    r.bold = True
            else:
                qty = e["qty"]
                if e.get("unit_cost") is not None:
                    qty = (qty + " × " if qty else "") + fmt_amount(e["unit_cost"])
                row[0].text, row[1].text, row[2].text, row[3].text = e["num"], e["name"], qty, fmt_amount(e["amount"])
        row = table.add_row().cells
        row[1].text, row[3].text = "ИТОГО", fmt_amount(struct["total"])
        for c in (1, 3):
            for r in row[c].paragraphs[0].runs:
                r.bold = True
    else:
        for line in budget_text.split("\n"):
            document.add_paragraph(re.sub(r"[*#`]", "", line).rstrip())
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
