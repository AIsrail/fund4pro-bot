"""Заполнение формы донора через явную СХЕМУ полей + структурированный
батч-вызов модели (JSON по стабильному field_id) — вместо нечёткого
сопоставления текста между отдельно сгенерированным markdown-документом и
ячейками шаблона (см. docx_template_fill.py, оставлен как fallback).

ПОЧЕМУ ПЕРЕДЕЛАНО (архитектурный аудит после серии багов за одну неделю:
заявка не по форме, дублирующиеся ответы в разных полях, потерянные разделы,
XYZ вместо реальных данных): старый пайплайн работал в два физически не
связанных шага — (1) generate_final_document пишет ПОЛНЫЙ текст заявки
свободным markdown, ничего не зная про реальные ячейки шаблона; (2)
docx_template_fill.py пытается угадать, какой кусок этого текста куда
вписать, через нечёткое сравнение текста вопроса с текстом заголовка/метки.
Это и есть источник большинства багов — угадывание неизбежно иногда
промахивается, задваивает совпадения или не находит вообще ничего.

Новый подход убирает угадывание целиком: сначала СЧИТЫВАЕМ реальную
структуру шаблона напрямую (какие ячейки где, с каким вопросом) — это
не угадывание, а точное чтение. Затем ОДНИМ (или несколькими батчами для
длинных форм) структурированным вызовом просим модель ответить по стабильному
field_id на каждый вопрос сразу с учётом всего проекта. Ответ приходит уже
привязанным к конкретной ячейке — записываем напрямую, сопоставлять по
тексту не нужно вообще.

Область применения: только табличная часть шаблона (там же, где раньше
работал fuzzy-matching — все содержательные поля реальных форм доноров на
практике лежат в таблицах). Чекбоксы/флажки вне таблиц (обычные параграфы) —
как и раньше, эта функция их не трогает; известное, некритичное ограничение,
не регрессия относительно старого поведения (там их тоже не было)."""

import copy
import logging
import re
from dataclasses import dataclass
from typing import Any

import docx
from docx.shared import Pt
from docx.table import _Row

from xyz_highlight import add_text_xyz_highlighted

logger = logging.getLogger("fund4pro.docx_schema_fill")

PROSE_CELL_MIN_CHARS = 300  # ячейка с вложенной таблицей и вопросом длиннее — есть и текстовая часть
CHUNK_SIZE = 12  # полей на один структурированный вызов — держит ответ в разумных max_tokens


@dataclass
class FieldSpec:
    field_id: str
    question: str
    kind: str  # "kv" (короткий факт рядом с меткой) | "section" (открытый абзац) | "table" (вложенная таблица)
    target: Any  # "kv"/"section": docx.table._Cell; "table": (nested_table, template_row_idx, total_row_idx)
    columns: list[str] | None = None  # только для kind == "table" — заголовки колонок по порядку
    before_table: bool = False  # "section" в ячейке с вложенной таблицей: текст идёт ВЫШЕ таблицы


def extract_template_schema(doc: "docx.Document") -> list[FieldSpec]:
    """Читает РЕАЛЬНУЮ структуру шаблона напрямую из таблиц — тот же обход,
    что раньше использовался только для чтения меток под fuzzy-matching,
    теперь ещё и запоминает прямую ссылку на целевую ячейку для каждого
    поля, так что заполнение потом не требует повторного поиска."""
    fields: list[FieldSpec] = []
    counter = 0

    # Снимаем свойства плавающей таблицы — как в старом пайплайне, иначе
    # Word может неправильно спозиционировать таблицу после правки текста.
    for table in doc.tables:
        tblpPr = table._tbl.tblPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblpPr")
        if tblpPr is not None:
            table._tbl.tblPr.remove(tblpPr)

    for table in doc.tables:
        for row in table.rows:
            unique_cells = []
            for c in row.cells:
                if not any(c._tc == u._tc for u in unique_cells):
                    unique_cells.append(c)

            if len(unique_cells) >= 2:
                # Метка -> соседняя ПУСТАЯ ячейка = поле ввода для этой метки
                for j in range(len(unique_cells)):
                    lbl = unique_cells[j].text.strip()
                    if not lbl:
                        continue
                    if j + 1 < len(unique_cells) and not unique_cells[j + 1].text.strip():
                        counter += 1
                        fields.append(FieldSpec(f"f{counter}", lbl, "kv", unique_cells[j + 1]))
            elif len(unique_cells) == 1:
                cell = unique_cells[0]
                text = cell.text.strip()
                if not text:
                    continue
                # РЕАЛЬНЫЙ ИНЦИДЕНТ: некоторые вопросы формы (план мероприятий,
                # построчный бюджет) требуют не абзац текста, а ВЛОЖЕННУЮ В
                # ЯЧЕЙКУ ТАБЛИЦУ — донор жёстко рассчитывает на построчный
                # формат. cell.text не включает текст вложенных таблиц (он
                # читается отдельно через cell.tables), так что раньше такие
                # поля тихо классифицировались как "section" и вложенная
                # таблица оставалась пустой навсегда — абзац дописывался
                # рядом, а не в предназначенные для этого строки.
                if cell.tables:
                    nested = cell.tables[0]
                    header, template_row_idx, total_row_idx = _find_nested_table_rows(nested)
                    if header and template_row_idx is not None:
                        counter += 1
                        fields.append(FieldSpec(
                            f"f{counter}", text, "table",
                            (nested, template_row_idx, total_row_idx), columns=header,
                        ))
                        # РЕАЛЬНЫЙ ИНЦИДЕНТ: ячейка «ПРОЕКТ: цель, задачи, деятельность,
                        # результат... + рабочий план» содержит и текстовые вопросы, и
                        # вложенную таблицу. Раньше поле считалось ТОЛЬКО таблицей —
                        # описание проекта (главный раздел заявки) не писалось вообще.
                        # Длинная формулировка = есть текстовая часть -> отдельное
                        # section-поле, ответ вписывается выше таблицы.
                        if len(text) > PROSE_CELL_MIN_CHARS:
                            counter += 1
                            fields.append(FieldSpec(
                                f"f{counter}", "[только текстовая часть] " + text, "section",
                                cell, before_table=True,
                            ))
                        continue
                    # Вложенная таблица есть, но без узнаваемого шаблона строки
                    # (нет ни одной полностью пустой строки после заголовка) —
                    # падаем обратно на section, чтобы хотя бы текст не потерять.
                counter += 1
                fields.append(FieldSpec(f"f{counter}", text, "section", cell))

    return fields


def _find_nested_table_rows(nested_table) -> tuple[list[str], int | None, int | None]:
    """Определяет заголовок вложенной таблицы, индекс пустой строки-шаблона
    (куда вписывать реальные данные) и, если есть, индекс строки "ИТОГО"
    (для итоговой суммы бюджета)."""
    if len(nested_table.rows) < 2:
        return [], None, None
    header = [c.text.strip() for c in nested_table.rows[0].cells]
    template_row_idx = None
    total_row_idx = None
    for ri in range(1, len(nested_table.rows)):
        cells_text = [c.text.strip() for c in nested_table.rows[ri].cells]
        if any(t.upper() in ("ИТОГО", "ИТОГ", "TOTAL", "ВСЕГО") for t in cells_text if t):
            total_row_idx = ri
        elif template_row_idx is None and all(t == "" for t in cells_text):
            template_row_idx = ri
    return header, template_row_idx, total_row_idx


async def build_field_answers(fields: list[FieldSpec], session_data: dict) -> dict[str, str]:
    """Отвечает на все поля схемы КРОМЕ "table" батчами по CHUNK_SIZE — один
    структурированный JSON-вызов на батч (llm.fill_form_fields_batch), не
    нечёткое сопоставление. "table"-поля отвечаются отдельно, через
    build_table_answers — форма их ответа другая (массив строк, а не одна
    строка на field_id), так что в общий батч они не годятся.

    Передаёт уже полученные ответы в контекст следующего батча, чтобы цифры
    (сумма бюджета, даты) оставались согласованными между батчами."""
    from llm import fill_form_fields_batch

    flat_fields = [f for f in fields if f.kind != "table"]
    answers: dict[str, str] = {}
    for i in range(0, len(flat_fields), CHUNK_SIZE):
        chunk = flat_fields[i:i + CHUNK_SIZE]
        chunk_dicts = [{"field_id": f.field_id, "question": f.question, "kind": f.kind} for f in chunk]
        try:
            result = await fill_form_fields_batch(chunk_dicts, session_data, already_answered=answers)
        except Exception:
            logger.warning("build_field_answers: batch %d-%d failed, continuing with rest", i, i + len(chunk), exc_info=True)
            result = {}
        answers.update(result)
    return answers


async def build_table_answers(fields: list[FieldSpec], session_data: dict) -> dict[str, list[list[str]]]:
    """Отдельный проход для "table"-полей (план мероприятий, построчный
    бюджет) — по одному вызову на поле, т.к. таких полей обычно 1-2 на форму
    и их ответ структурно другой (массив строк по колонкам)."""
    from llm import fill_table_field

    answers: dict[str, list[list[str]]] = {}
    for f in fields:
        if f.kind != "table":
            continue
        try:
            rows = await fill_table_field(f.question, f.columns or [], session_data)
        except Exception:
            logger.warning("build_table_answers: field %s failed", f.field_id, exc_info=True)
            rows = []
        if rows:
            _nudge_round_total(rows, f.columns)
            answers[f.field_id] = rows
    return answers


XYZ_MAX_SHARE = 0.10  # не более 10% значений в заявке могут быть маркерами XYZ
_NUM_RE = re.compile(r"\d[\d\s.,]*")


def xyz_ratio(texts: list[str]) -> float:
    """Доля XYZ среди всех "значений" (числа + XYZ) в текстах заявки."""
    xyz = sum(t.count("XYZ") for t in texts)
    nums = sum(len(_NUM_RE.findall(t.replace("XYZ", " "))) for t in texts)
    total = xyz + nums
    return xyz / total if total else 0.0


async def enforce_xyz_budget(
    answers: dict[str, str], table_answers: dict[str, list[list[str]]], session_data: dict,
) -> float:
    """Если XYZ больше 10% — один проход llm.resolve_placeholders_batch по
    фрагментам с XYZ (замена на оценки от бюджета). Правит answers и
    table_answers на месте, возвращает итоговую долю XYZ."""
    from llm import resolve_placeholders_batch

    def all_texts() -> list[str]:
        return list(answers.values()) + [c for rows in table_answers.values() for r in rows for c in r]

    ratio = xyz_ratio(all_texts())
    if ratio <= XYZ_MAX_SHARE:
        return ratio

    items = [{"id": fid, "text": t} for fid, t in answers.items() if "XYZ" in t]
    for fid, rows in table_answers.items():
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                if "XYZ" in cell:
                    items.append({"id": f"{fid}|{ri}|{ci}", "text": cell})

    for i in range(0, len(items), 12):
        resolved = await resolve_placeholders_batch(items[i:i + 12], session_data)
        for key, new_text in resolved.items():
            if "|" in key:
                fid, ri, ci = key.split("|")
                try:
                    table_answers[fid][int(ri)][int(ci)] = new_text
                except (KeyError, IndexError, ValueError):
                    pass
            elif key in answers:
                answers[key] = new_text

    final = xyz_ratio(all_texts())
    logger.info("enforce_xyz_budget: XYZ share %.0f%% -> %.0f%%", ratio * 100, final * 100)
    return final


_COUNTRY_OR_YEAR_RE = re.compile(
    r"^(19|20)\d{2}$|^(германия|россия|казахстан|кыргызстан|узбекистан|таджикистан|сша|"
    r"germany|usa|kyrgyzstan|kazakhstan|uzbekistan|tajikistan)$", re.IGNORECASE)


def sanitize_contact_answers(fields: list[FieldSpec], answers: dict[str, str]) -> int:
    """Страховка от «правдоподобного мусора» в контактах (реальный инцидент:
    «Контактное лицо: Германия», сайт/телефон/email: «2026»). Если значение
    по форме не может быть телефоном/email/сайтом/ФИО — заменяем на XYZ,
    чтобы пользователь увидел красное поле, а не поверил выдумке.
    Возвращает число исправленных полей."""
    fixed = 0
    for f in fields:
        if f.kind != "kv":
            continue
        val = (answers.get(f.field_id) or "").strip()
        if not val or val.upper() == "XYZ":
            continue
        label = f.question.lower()
        bad = False
        if "почт" in label or "e-mail" in label or "email" in label:
            bad = "@" not in val
        elif "телефон" in label or label.strip() in ("tel", "phone"):
            bad = sum(ch.isdigit() for ch in val) < 6
        elif "сайт" in label or "соцсет" in label or "website" in label:
            bad = not re.search(r"[.@]|instagram|facebook|telegram", val, re.IGNORECASE)
        elif label.startswith("контактное лицо") or "contact person" in label:
            words = [w for w in re.split(r"\s+", val) if re.search(r"[^\W\d_]", w)]
            bad = len(words) < 2 or bool(_COUNTRY_OR_YEAR_RE.match(val))
        if bad:
            logger.warning("sanitize_contact_answers: %r -> XYZ (label=%r)", val[:60], f.question[:60])
            answers[f.field_id] = "XYZ"
            fixed += 1
    return fixed


def _fill_kv_cell(target, val: str) -> None:
    target.text = ""
    p = target.paragraphs[0] if target.paragraphs else target.add_paragraph()
    add_text_xyz_highlighted(p, val, font_name="Times New Roman", font_size=Pt(10.5))


def _append_section_answer(cell, content: str, before_table: bool = False) -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    add_text_xyz_highlighted(p, f"\n{content}", font_name="Times New Roman", font_size=Pt(11))
    if before_table and cell.tables:
        # add_paragraph() кладёт абзац в конец ячейки (под вложенную таблицу) —
        # переносим выше неё; если прямо над таблицей короткий заголовок
        # («Рабочий план реализации проекта»), текст встаёт над этим заголовком.
        tbl = cell.tables[0]._tbl
        anchor = tbl
        prev = tbl.getprevious()
        while prev is not None and prev.tag.endswith("}p") and not "".join(prev.itertext()).strip():
            anchor = prev
            prev = prev.getprevious()
        if prev is not None and prev.tag.endswith("}p") and len("".join(prev.itertext()).strip()) < 80:
            anchor = prev
        anchor.addprevious(p._p)


_NUMBER_RE = re.compile(r"[\d\s]+(?:[.,]\d+)?")
_MONEY_COL_RE = re.compile(r"стоимост|сумма|amount|cost|usd|\$", re.IGNORECASE)


def _nudge_round_total(rows: list[list[str]], columns: list[str] | None) -> None:
    """РЕАЛЬНЫЙ ИНЦИДЕНТ (жалоба пользователя): итоговая сумма бюджета вышла
    ровно 5000 — для донора это красный флаг «заявитель не считал, подогнал
    под лимит» (это правило уже было в SYSTEM_PROMPT текстом, но слабо
    соблюдается, когда строки бюджета — это проценты от лимита: админ 12% +
    M&E 6% + контингенси 4% от РОВНОЙ суммы гранта механически дают ровный
    итог). Промпт-инструкцию усилили (см. fill_table_field в llm.py), но
    результат зависит от модели — здесь код детерминированно подстраховывает:
    если сумма всё равно кратна 100, чуть меняет САМУЮ КРУПНУЮ строку, чтобы
    итог выглядел как реальный расчёт (…5140/…5250), а не подгонка под лимit.
    Мутирует rows на месте; молча ничего не делает, если это не денежная
    колонка (план мероприятий) или сумма и так не подозрительно круглая."""
    if not columns or not _MONEY_COL_RE.search(columns[-1]):
        return
    valid = [(i, v) for i, r in enumerate(rows) if (v := _extract_number(r[-1])) is not None and v > 0]
    if len(valid) < 2:
        return
    total = sum(v for _, v in valid)
    if total <= 0 or total % 100 != 0:
        return  # не подозрительно круглая — не трогаем реальный ответ модели
    idx, biggest = max(valid, key=lambda iv: iv[1])
    # Детерминированная (не одинаковая для любой суммы) добавка ~1-5% от
    # крупнейшей строки, диапазон 20-260 — по составу самих строк, не по total.
    seed = sum(len(c) for r in rows for c in r) + len(rows)
    delta = 20 + (seed * 37) % 240
    new_val = biggest + delta
    original = rows[idx][-1]
    prefix = "$" if "$" in original else ""
    formatted = f"{new_val:,.0f}".replace(",", " ") if new_val >= 1000 else f"{new_val:.0f}"
    rows[idx][-1] = f"{prefix}{formatted}"


def _extract_number(text: str) -> float | None:
    """Best-effort: вытаскивает первое число из строки вида '1 820 USD' —
    используется только чтобы посчитать ИТОГО по бюджетной таблице; если
    формат неожиданный, просто пропускаем строку в сумме, не роняя запись."""
    m = _NUMBER_RE.search(text)
    if not m:
        return None
    raw = m.group(0).replace(" ", "").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return None


def _duplicate_row_after(nested_table, anchor_tr, template_tr) -> "_Row":
    """Клонирует XML строки-шаблона и вставляет копию сразу ПОСЛЕ anchor_tr
    (не всегда после самого шаблона!) — python-docx не даёт готового API
    "вставить строку между другими", это стандартный обходной путь через
    прямую работу с lxml-элементом.

    РЕАЛЬНЫЙ ИНЦИДЕНТ (поймано локальным тестом на реалистичных данных, ещё
    до релиза): если каждую новую строку вставлять "сразу после
    template_tr", то при N > 1 строках получается LIFO-порядок — последняя
    добавленная строка оказывается САМОЙ ПЕРВОЙ (3, 2, 1 вместо 1, 2, 3).
    Вызывающий код обязан двигать anchor_tr вперёд на каждую новую строку."""
    new_tr = copy.deepcopy(template_tr)
    anchor_tr.addnext(new_tr)
    return _Row(new_tr, nested_table)


def _fill_row_cells(row: "_Row", values: list[str]) -> None:
    for cell, val in zip(row.cells, values):
        cell.text = ""
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        add_text_xyz_highlighted(p, val, font_name="Times New Roman", font_size=Pt(10))


def _fill_table_field(field: FieldSpec, rows: list[list[str]]) -> bool:
    """Вписывает построчные данные во вложенную таблицу вместо одной пустой
    строки-шаблона: клонирует строку-шаблон под КАЖДУЮ реальную строку
    ответа, затем удаляет исходный пустой шаблон. Если у таблицы есть строка
    "ИТОГО" — досчитывает сумму по последней колонке (обычно "Стоимость"/
    "Сумма") и вписывает её в последнюю ячейку этой строки."""
    nested_table, template_row_idx, total_row_idx = field.target
    if not rows:
        return False

    template_tr = nested_table.rows[template_row_idx]._tr
    anchor_tr = template_tr
    for values in rows:
        row_obj = _duplicate_row_after(nested_table, anchor_tr, template_tr)
        _fill_row_cells(row_obj, values)
        anchor_tr = row_obj._tr  # следующая строка встаёт сразу после ЭТОЙ, не после шаблона

    # Исходная пустая строка-шаблон больше не нужна — удаляем её XML-узел.
    template_tr.getparent().remove(template_tr)

    if total_row_idx is not None:
        # total_row_idx был вычислен ДО клонирования/удаления — строки таблицы
        # сдвинулись на len(rows) - 1 (клоны добавлены, шаблон удалён).
        try:
            total_row = nested_table.rows[total_row_idx + len(rows) - 1]
            total = sum(v for v in (_extract_number(r[-1]) for r in rows) if v is not None)
            if total:
                last_cell = total_row.cells[-1]
                last_cell.text = ""
                p = last_cell.paragraphs[0] if last_cell.paragraphs else last_cell.add_paragraph()
                run = p.add_run(f"{total:,.0f}".replace(",", " "))
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
        except IndexError:
            logger.warning("_fill_table_field: could not locate ИТОГО row after cloning")
    return True


def write_answers_to_template(
    fields: list[FieldSpec], answers: dict[str, str], table_answers: dict[str, list[list[str]]] | None = None,
) -> tuple[int, int, int]:
    """Прямая, детерминированная запись — каждый field.target уже указывает
    ровно на ту ячейку (или вложенную таблицу), куда должен попасть его
    ответ. Сопоставлять по тексту не нужно: связь установлена один раз при
    извлечении схемы."""
    table_answers = table_answers or {}
    filled_kv = 0
    filled_sections = 0
    filled_tables = 0
    for f in fields:
        if f.kind == "table":
            rows = table_answers.get(f.field_id)
            if rows and _fill_table_field(f, rows):
                filled_tables += 1
            continue
        val = (answers.get(f.field_id) or "").strip()
        if not val:
            continue
        if f.kind == "kv":
            _fill_kv_cell(f.target, val)
            filled_kv += 1
        else:
            if val not in f.target.text:
                _append_section_answer(f.target, val, before_table=f.before_table)
                filled_sections += 1
    return filled_kv, filled_sections, filled_tables


async def fill_donor_docx_template_v2(template_path: str, output_path: str, session: dict) -> tuple[bool, str]:
    """Новый основной путь заполнения формы донора. Возвращает
    (успех, текст_для_проверки) — текст_для_проверки — это все реально
    записанные значения полей, склеенные вместе; используется вызывающим
    кодом (agent_engine.py) для того же самого подсчёта XYZ-плейсхолдеров,
    что и раньше, без дублирования этой логики здесь."""
    try:
        doc = docx.Document(template_path)
    except Exception as e:
        logger.warning("fill_donor_docx_template_v2: could not open template %s: %s", template_path, e)
        return False, ""

    fields = extract_template_schema(doc)
    if not fields:
        logger.warning("fill_donor_docx_template_v2: no fillable fields found in %s", template_path)
        return False, ""

    answers = await build_field_answers(fields, session)
    sanitize_contact_answers(fields, answers)
    table_answers = await build_table_answers(fields, session)
    try:
        await enforce_xyz_budget(answers, table_answers, session)
    except Exception:
        logger.warning("enforce_xyz_budget failed, continuing with unresolved XYZ", exc_info=True)
    filled_kv, filled_sections, filled_tables = write_answers_to_template(fields, answers, table_answers)

    logger.info(
        "fill_donor_docx_template_v2: filled %d/%d kv fields, %d sections, %d tables (of %d total fields) in %s",
        filled_kv, sum(1 for f in fields if f.kind == "kv"), filled_sections, filled_tables, len(fields), template_path,
    )

    # Та же защита, что была в старом пайплайне: форма с солидным числом
    # таблиц, но почти ничем не заполненная — считаем неуспехом, а не тихо
    # сохраняем полупустой файл.
    if len(doc.tables) >= 5 and (filled_kv < 5 or filled_sections < 1):
        logger.warning(
            "fill_donor_docx_template_v2: only %d kv + %d sections filled — rejecting partial fill",
            filled_kv, filled_sections,
        )
        return False, ""

    if filled_kv + filled_sections + filled_tables == 0:
        return False, ""

    doc.save(output_path)
    table_text = "\n".join(
        " | ".join(cell for row in rows for cell in row) for rows in table_answers.values()
    )
    text_for_check = "\n".join(answers.get(f.field_id, "") for f in fields) + "\n" + table_text
    return True, text_for_check
