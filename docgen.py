"""Конвертер markdown-текста в .docx, понимающий реальный markdown, а не
только заголовки/буллеты: таблицы (| ... | ... |), инлайн **жирный** и
*курсив* текст. Раньше таблицы и жирный текст выводились как сырые символы
(`|---|---|---|`, `**слово**`) прямо в тексте документа — это и есть баг
"много ненужных символов", который видел пользователь в финальном .docx.
"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from xyz_highlight import add_text_xyz_highlighted

_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_INLINE_ITALIC_RE = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")


def markdown_to_docx(text: str, title: str | None, path: str) -> str:
    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    lines = text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        # Markdown-таблица: строка с | ..., затем строка-разделитель
        # |---|---|---|, затем строки данных — пока встречаются строки с |.
        if stripped.startswith("|") and i + 1 < n and _is_table_separator(lines[i + 1]):
            table_rows, consumed = _parse_table_block(lines[i:])
            _render_table(doc, table_rows)
            i += consumed
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, stripped[2:].strip())
        else:
            p = doc.add_paragraph()
            _add_inline_runs(p, stripped)
        i += 1

    doc.save(path)
    return path


def _is_table_separator(line: str) -> bool:
    s = line.strip()
    if not s.startswith("|"):
        return False
    inner = s.strip("|")
    cells = [c.strip() for c in inner.split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{1,}:?", c) for c in cells if c)


def _parse_table_block(lines: list[str]) -> tuple[list[list[str]], int]:
    """lines[0] = header row, lines[1] = separator, lines[2:] = data rows
    (пока встречаются строки, начинающиеся с |). Возвращает (строки таблицы
    как список списков ячеек, количество потреблённых строк исходного
    текста)."""
    header = _split_table_row(lines[0])
    rows = [header]
    consumed = 2  # header + separator
    for line in lines[2:]:
        if not line.strip().startswith("|"):
            break
        rows.append(_split_table_row(line))
        consumed += 1
    return rows, consumed


def _split_table_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    return [c.strip() for c in inner.split("|")]


def _render_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline_runs(p, cell_text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()  # отступ после таблицы


def _add_inline_runs(paragraph, text: str) -> None:
    """Разбивает текст на обычные/жирные/курсивные фрагменты по **/*
    маркерам markdown и добавляет их как отдельные runs — вместо того,
    чтобы звёздочки просто попадали в текст как есть. Любое вхождение
    XYZ-плейсхолдера внутри любого фрагмента красится красным жирным (см.
    xyz_highlight) — так пользователь физически видит на вычитке, что
    нужно заменить перед подачей, а не выискивает их в сплошном тексте."""
    pos = 0
    tokens = []
    for m in _INLINE_BOLD_RE.finditer(text):
        if m.start() > pos:
            tokens.append(("normal", text[pos:m.start()]))
        tokens.append(("bold", m.group(1)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("normal", text[pos:]))
    if not tokens:
        tokens = [("normal", text)]

    for kind, chunk in tokens:
        if kind == "bold":
            add_text_xyz_highlighted(paragraph, chunk, bold=True)
            continue
        # внутри обычного фрагмента ещё может быть *курсив*
        sub_pos = 0
        for m in _INLINE_ITALIC_RE.finditer(chunk):
            if m.start() > sub_pos:
                add_text_xyz_highlighted(paragraph, chunk[sub_pos:m.start()])
            add_text_xyz_highlighted(paragraph, m.group(1), italic=True)
            sub_pos = m.end()
        if sub_pos < len(chunk):
            add_text_xyz_highlighted(paragraph, chunk[sub_pos:])
