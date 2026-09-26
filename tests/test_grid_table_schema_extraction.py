"""Регрессионный тест для extract_template_schema — ГРИД-таблицы (строка
заголовка именует каждую колонку: план мероприятий, риски, построчный
бюджет).

РЕАЛЬНАЯ ЖАЛОБА (живой тест, донор FAO, 26.09.2026): "бот теперь работает
только с шаблоном, даже уже лучше находит данные, но упорно не заполняет
его до конца — часть разделов пустые". Проверка присланного файла показала:
для строк вида [метка, пусто, пусто, пусто] старый алгоритм создавал РОВНО
ОДНО поле на строку — метку в col[0] в пару с ПЕРВОЙ соседней пустой
ячейкой. Остальные пустые ячейки той же строки НИКОГДА не становились
полями вообще — ни на каком провайдере, потому что вопрос для них не
существовал в схеме в принципе. План мероприятий заполнялся только в
колонке "Мероприятие" ("Сроки"/"Ответственный" — пустые НАВСЕГДА), а
таблицы риск/бюджет без родной метки в col[0] не заполнялись НИ ОДНИМ
полем. Это НЕ проблема качества модели — вопросы для этих ячеек не
задавались вообще, слабый или сильный провайдер тут не при чём.

    python -m tests.test_grid_table_schema_extraction
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx


def _add_grid_table(doc, header: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    for ci, h in enumerate(header):
        table.rows[0].cells[ci].text = h
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    return table


def test_activity_plan_grid_gets_a_field_for_every_empty_column_not_just_the_first():
    """План мероприятий: [номер, пусто, пусто, пусто] должен дать ТРИ поля
    на строку (Мероприятие/Сроки/Ответственный), не одно."""
    from docx_schema_fill import extract_template_schema

    doc = docx.Document()
    _add_grid_table(
        doc,
        ["№", "Мероприятие", "Сроки / Период", "Ответственный"],
        [["1.1.", "", "", ""], ["1.2.", "", "", ""]],
    )
    fields = extract_template_schema(doc)
    by_question = {f.question: f for f in fields}

    assert "Мероприятие — «1.1.»" in by_question
    assert "Сроки / Период — «1.1.»" in by_question
    assert "Ответственный — «1.1.»" in by_question
    assert "Мероприятие — «1.2.»" in by_question
    assert "Сроки / Период — «1.2.»" in by_question
    assert "Ответственный — «1.2.»" in by_question
    assert len(fields) == 6, f"expected exactly 6 fields (3 columns x 2 rows), got {len(fields)}"
    print("OK: an activity-plan grid gets a field for every empty column in every row, not just the first")


def test_unlabeled_small_grid_like_risks_gets_fully_extended():
    """Таблица рисков БЕЗ номеров строк (нет метки в col[0] вообще) —
    должна получить поля на КАЖДУЮ ячейку каждой строки, раз таблица
    маленькая (явный чек-лист 'заполни все')."""
    from docx_schema_fill import extract_template_schema

    doc = docx.Document()
    _add_grid_table(
        doc,
        ["Риски", "Вероятность\n(Н, С, В)", "Меры по их снижению или решению"],
        [["", "", ""], ["", "", ""], ["", "", ""]],
    )
    fields = extract_template_schema(doc)
    questions = [f.question for f in fields]

    assert len(fields) == 9, f"expected 3 columns x 3 unlabeled rows = 9 fields, got {len(fields)}: {questions}"
    assert any(q.startswith("Риски — ") for q in questions)
    assert any(q.startswith("Вероятность") for q in questions)
    assert any(q.startswith("Меры по их снижению") for q in questions)
    print("OK: a small unlabeled risk-style grid gets every cell in every row extracted as its own field")


def test_large_non_budget_table_with_many_unlabeled_blank_rows_is_not_exploded():
    """Защита от противоположной крайности: НЕ-бюджетная таблица с ДЕСЯТКАМИ
    пустых безымянных под-строк (растущий список 'добавь свои позиции при
    необходимости') НЕ должна превратиться в десятки придуманных строк —
    только именованные строки достраиваются. Бюджетные таблицы — исключение,
    см. test_budget_grid_always_gets_itemized_continuation_rows ниже."""
    from docx_schema_fill import extract_template_schema

    doc = docx.Document()
    header = ["Дата", "Событие", "Участники", "Заметки"]
    rows = [["Запись 1", "", "", ""]] + [["", "", "", ""] for _ in range(20)]
    _add_grid_table(doc, header, rows)

    fields = extract_template_schema(doc)
    questions = [f.question for f in fields]

    # Именованная строка "Запись 1" достроена (3 пустые колонки).
    assert sum(1 for q in questions if "«Запись 1»" in q) == 3, questions
    # Ни одно поле не создано для 20 безымянных пустых строк большой таблицы.
    assert not any("пункт" in q for q in questions), (
        f"unlabeled blank rows in a large (>6 rows) NON-budget table must NOT be auto-extended, "
        f"got: {questions}"
    )
    print("OK: a large non-budget table's many unlabeled blank rows are left alone — only named rows get extended")


def test_budget_grid_always_gets_itemized_continuation_rows_regardless_of_table_size():
    """РЕАЛЬНАЯ ПРОСЬБА ВЛАДЕЛЬЦА (26.09.2026): "бюджет надо всегда писать
    детально... никогда не обобщайте по категориям". Пустые строки-продолжения
    ПОД категорией в построчном бюджете должны доставаться в поля независимо
    от размера таблицы (в отличие от обычных нет-бюджетных таблиц выше) —
    донор сам ограничивает их число своим шаблоном, плодить бесконечно нечем."""
    from docx_schema_fill import extract_template_schema

    doc = docx.Document()
    header = ["Статьи", "Ед.", "Стоимость за ед.", "Количество", "ВСЕГО"]
    rows = (
        [["Оборудование", "", "", "", ""]] + [["", "", "", "", ""] for _ in range(3)]
        + [["Публикации", "", "", "", ""]] + [["", "", "", "", ""] for _ in range(2)]
    )
    _add_grid_table(doc, header, rows)

    fields = extract_template_schema(doc)
    questions = [f.question for f in fields]

    # "Оборудование": 1 категорийная строка + 3 пустых продолжения = 4 строки,
    # по 4 обычных поля на строку (Ед./Стоимость/Количество/ВСЕГО) + 1
    # "конкретная позиция" поле для col[0] в каждой из 3 строк-продолжений.
    assert sum(1 for q in questions if "«Оборудование»" in q and "конкретная позиция" not in q) == 16, questions
    equipment_continuation = [q for q in questions if "Оборудование" in q and "конкретная позиция" in q]
    assert len(equipment_continuation) == 3, (
        f"expected one 'Статьи — конкретная позиция...' field per continuation row under "
        f"'Оборудование' (3 rows), got {len(equipment_continuation)}: {questions}"
    )
    # "Публикации": 1 категорийная строка + 2 пустых продолжения = 3 строки x 4 = 12.
    assert sum(1 for q in questions if "«Публикации»" in q and "конкретная позиция" not in q) == 12, questions
    publications_continuation = [q for q in questions if "Публикации" in q and "конкретная позиция" in q]
    assert len(publications_continuation) == 2, questions
    print("OK: a budget grid's blank continuation rows under each category are always itemized, any table size")


def test_non_grid_label_value_table_is_unaffected():
    """Обычная форма профиля организации (заголовок НЕ у всех ячеек
    непустой — первая строка САМА является парой метка:значение) должна
    заполняться старой логикой без изменений."""
    from docx_schema_fill import extract_template_schema

    doc = docx.Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Название организации:"
    table.rows[0].cells[1].text = ""
    table.rows[1].cells[0].text = "Email:"
    table.rows[1].cells[1].text = ""
    table.rows[2].cells[0].text = "Телефон:"
    table.rows[2].cells[1].text = ""

    fields = extract_template_schema(doc)
    questions = {f.question for f in fields}
    assert questions == {"Название организации:", "Email:", "Телефон:"}, questions
    print("OK: a plain label:value profile table is untouched by the grid-extraction logic")


if __name__ == "__main__":
    test_activity_plan_grid_gets_a_field_for_every_empty_column_not_just_the_first()
    test_unlabeled_small_grid_like_risks_gets_fully_extended()
    test_large_non_budget_table_with_many_unlabeled_blank_rows_is_not_exploded()
    test_budget_grid_always_gets_itemized_continuation_rows_regardless_of_table_size()
    test_non_grid_label_value_table_is_unaffected()
    print("\nAll grid-table-schema-extraction tests passed.")
