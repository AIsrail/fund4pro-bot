"""Юнит-тесты правил качества (quality_rules.py) и жёлтой подсветки пометок.

    python -m tests.test_quality_rules
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

import agent_roadmap
import llm
import quality_rules
from xyz_highlight import add_text_xyz_highlighted


def test_admin_rule_not_fixed_percent_and_in_prompts():
    r = quality_rules.ADMIN_SHARE_RULE
    assert "НЕ ФИКСИРОВАННЫЙ" in r and "попроси пользователя" in r
    assert quality_rules.ADMIN_SHARE_RULE in llm.SYSTEM_PROMPT
    assert quality_rules.ADMIN_SHARE_RULE in agent_roadmap.FULL_ROADMAP_SYSTEM_PROMPT
    assert "≤10-20%" not in llm.SYSTEM_PROMPT


def test_applicant_comment_yellow_and_xyz_red():
    p = Document().add_paragraph()
    add_text_xyz_highlighted(
        p, "Сумма XYZ USD. [КОММЕНТАРИЙ ДЛЯ ЗАЯВИТЕЛЯ — НЕ ДЛЯ ДОНОРА: проверьте XYZ лимит] Конец."
    )
    runs = p.runs
    yellow = [r for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert len(yellow) == 1 and yellow[0].text.startswith("[КОММЕНТАРИЙ ДЛЯ ЗАЯВИТЕЛЯ")
    red = [r for r in runs if r.text == "XYZ" and r.font.color.rgb is not None]
    assert len(red) == 1  # XYZ внутри комментария отдельно не красится
    assert "".join(r.text for r in runs).endswith("Конец.")


def test_plain_text_without_comment_unchanged():
    p = Document().add_paragraph()
    add_text_xyz_highlighted(p, "Обычный XYZ текст")
    assert "".join(r.text for r in p.runs) == "Обычный XYZ текст"
    assert not any(r.font.highlight_color for r in p.runs)


if __name__ == "__main__":
    test_admin_rule_not_fixed_percent_and_in_prompts()
    test_applicant_comment_yellow_and_xyz_red()
    test_plain_text_without_comment_unchanged()
    print("All quality-rules tests passed.")
