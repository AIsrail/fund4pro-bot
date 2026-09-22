"""Заполнение бюджетного шаблона донора прямо в Excel (.xlsx), а не только
текстом в .docx.

Раньше даже если донор давал бюджет именно в Excel-форме, бот всё равно
выдавал только markdown-текст, обёрнутый в .docx — а донор просит
заполнить конкретно его файл (с его столбцами/формулами/листами).

Подход: читаем структуру шаблона (список непустых ячеек с координатами
и текстом — это подписи полей/шапка таблицы), просим LLM вернуть, В
КАКИЕ именно ячейки (по координатам вроде "B7") нужно вписать какие
значения, основываясь на уже согласованном с пользователем бюджете
(budget_text) — и пишем значения именно туда, не трогая остальной файл
(формулы/форматирование/другие листы остаются как есть).
"""

import io
import json
import logging

from openpyxl import load_workbook

logger = logging.getLogger("fund4pro.excel_fill")

MAX_CELLS_FOR_PROMPT = 400  # не заливать промпт тысячами пустых/служебных ячеек


def extract_xlsx_structure(content: bytes) -> str:
    """Возвращает текстовое описание структуры шаблона: лист, координата
    ячейки, её текущее содержимое — только непустые ячейки, чтобы LLM
    видела, куда именно (в какие подписанные поля/строки таблицы) нужно
    вписывать бюджетные значения."""
    wb = load_workbook(io.BytesIO(content), data_only=False)
    lines = []
    count = 0
    for sheet in wb.worksheets:
        lines.append(f"[Лист: {sheet.title}]")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip():
                    lines.append(f"{cell.coordinate}: {cell.value}")
                    count += 1
                    if count >= MAX_CELLS_FOR_PROMPT:
                        lines.append("... (обрезано, шаблон больше)")
                        return "\n".join(lines)
    return "\n".join(lines)


def fill_xlsx_template(content: bytes, cell_values: dict) -> bytes:
    """cell_values: {"Лист1": {"C7": "12000", ...}, ...}. Пишет значения
    ТОЛЬКО в указанные ячейки поверх копии оригинального файла — формулы,
    форматирование, остальные листы/ячейки не трогаются."""
    wb = load_workbook(io.BytesIO(content), data_only=False)
    for sheet_name, cells in cell_values.items():
        if sheet_name not in wb.sheetnames:
            # LLM могла ошибиться с именем листа — пробуем единственный лист
            if len(wb.sheetnames) == 1:
                sheet_name_actual = wb.sheetnames[0]
            else:
                continue
        else:
            sheet_name_actual = sheet_name
        ws = wb[sheet_name_actual]
        for coord, value in cells.items():
            try:
                ws[coord] = value
            except Exception as exc:
                logger.warning("failed to set cell %s!%s: %s", sheet_name_actual, coord, exc)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def fill_donor_xlsx_template_v1(template_path: str, output_path: str, session: dict) -> tuple[bool, str]:
    """xlsx-аналог docx_schema_fill.fill_donor_docx_template_v2 / pdf_form_fill.
    fill_donor_pdf_template_v1 — тот же контракт (успех, текст_для_проверки),
    чтобы agent_docgen._try_schema_fill мог выбирать путь по расширению
    файла. РЕАЛЬНЫЙ ИНЦИДЕНТ: этот модуль существовал и работал, но был
    подключён только в СТАРОМ FSM-хендлере (handlers/budget.py), который
    текущая агентная архитектура (bot.py подключает только agent_router) не
    вызывает вообще — для донора с бюджетным Excel-шаблоном (например NED)
    бот раньше никогда не пытался заполнить именно его файл."""
    try:
        with open(template_path, "rb") as fh:
            content = fh.read()
    except Exception as e:
        logger.warning("fill_donor_xlsx_template_v1: could not read %s: %s", template_path, e)
        return False, ""

    project_data = session.get("project_data", {})
    budget_text = project_data.get("activities_and_budget", "")
    if not budget_text.strip():
        logger.info("fill_donor_xlsx_template_v1: no activities_and_budget yet, nothing to map")
        return False, ""

    structure = extract_xlsx_structure(content)
    if not structure.strip():
        logger.info("fill_donor_xlsx_template_v1: empty xlsx structure (no labelled cells) in %s", template_path)
        return False, ""

    mapping = await generate_budget_cell_mapping(structure, budget_text)
    if not mapping:
        logger.warning("fill_donor_xlsx_template_v1: could not determine cell mapping for %s", template_path)
        return False, ""

    try:
        filled = fill_xlsx_template(content, mapping)
    except Exception:
        logger.warning("fill_donor_xlsx_template_v1: fill_xlsx_template failed", exc_info=True)
        return False, ""

    with open(output_path, "wb") as fh:
        fh.write(filled)

    n_cells = sum(len(cells) for cells in mapping.values())
    logger.info("fill_donor_xlsx_template_v1: filled %d cell(s) across %d sheet(s) in %s", n_cells, len(mapping), template_path)
    text_for_check = "\n".join(f"{sheet}!{coord}: {val}" for sheet, cells in mapping.items() for coord, val in cells.items())
    return True, text_for_check


async def generate_budget_cell_mapping(structure_text: str, budget_text: str) -> dict:
    """LLM решает, в какие координаты ячеек вписать какие суммы, основываясь
    на согласованном с пользователем бюджете. Возвращает {} при любой
    ошибке разбора — вызывающий код должен считать это 'не смогли
    автозаполнить, прикладываем оригинал как есть', а не падать."""
    from llm import call_claude

    system_prompt = (
        "Перед тобой структура Excel-шаблона бюджета донора (координаты "
        "непустых ячеек и их текст — это подписи полей/шапка таблицы) и "
        "согласованный с пользователем бюджет проекта. Определи, в какие "
        "именно ячейки нужно вписать суммы/цифры (обычно это пустые ячейки "
        "рядом/под текстовыми подписями статей расходов), и верни СТРОГО "
        "JSON без каких-либо пояснений в формате:\n"
        '{"Название листа": {"C7": "12000", "C8": "5000"}}\n\n'
        "Пиши в ячейки только числа/суммы (без валютных символов, если в "
        "шаблоне уже есть колонка с валютой) — ориентируйся на формат "
        "соседних заполненных ячеек, если такие есть. Если не уверен, в "
        "какую именно ячейку что-то писать — не пиши, лучше пропустить "
        "поле, чем испортить не ту ячейку. Не трогай ячейки, которые уже "
        "содержат текст/формулы (Excel-формулы начинаются с '=')."
    )
    try:
        raw = await call_claude(
            system_prompt,
            f"Структура шаблона:\n{structure_text}\n\nСогласованный бюджет:\n{budget_text}",
            max_tokens=2000,
        )
    except Exception as exc:
        logger.warning("generate_budget_cell_mapping failed: %s: %s", type(exc).__name__, exc)
        return {}
    if not raw:
        return {}
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end == -1:
        return {}
    try:
        return json.loads(raw[start:end + 1])
    except Exception as exc:
        logger.warning("generate_budget_cell_mapping: bad JSON: %s", exc)
        return {}
